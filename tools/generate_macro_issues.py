#!/usr/bin/env python3
"""
Generate GitHub Issue markdown bodies for the 4 Semantic Macro Issues
from the reconciled Semantic_Project_Health_Ledger.xlsx.
"""

import os
import openpyxl
from collections import defaultdict

REPO_ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
LEDGER_PATH = os.path.join(REPO_ROOT, "reports", "security", "Semantic_Project_Health_Ledger.xlsx")
SCRATCH_DIR = os.path.join(REPO_ROOT, "scratch")
os.makedirs(SCRATCH_DIR, exist_ok=True)

wb = openpyxl.load_workbook(LEDGER_PATH, data_only=True)

# 1. Root causes
ws2 = wb['02_Root_Causes']
h2 = [c for c in next(ws2.iter_rows(values_only=True))]
m2 = {h: i for i, h in enumerate(h2)}
rc_data = {}
for r in ws2.iter_rows(min_row=2, values_only=True):
    rc_data[r[m2['Root_Cause_ID']]] = {
        'title': r[m2['Root_Cause_Title']],
        'area': r[m2['Area']],
        'desc': r[m2['Description']],
    }

# 2. Remediation queue
ws7 = wb['07_Remediation_Queue']
h7 = [c for c in next(ws7.iter_rows(values_only=True))]
m7 = {h: i for i, h in enumerate(h7)}
rem_data = {}
for r in ws7.iter_rows(min_row=2, values_only=True):
    rem_data[r[m7['Queue_ID']]] = {
        'rc_id': r[m7['Root_Cause_ID']],
        'area': r[m7['Area']],
        'priority': r[m7['Priority']],
        'risk_tier': r[m7['Risk_Tier']],
        'deps': r[m7['Dependencies']],
        'suggested_scope': r[m7['Suggested_PR_Scope']],
    }

# 3. Active findings
ws1 = wb['01_Findings']
h1 = [c for c in next(ws1.iter_rows(values_only=True))]
m1 = {h: i for i, h in enumerate(h1)}
active_findings_by_rc = defaultdict(list)
for r in ws1.iter_rows(min_row=2, values_only=True):
    st = r[m1['Current_Status']]
    if st in ('STILL_PRESENT', 'PARTIALLY_FIXED'):
        fid = r[m1['Finding_ID']]
        rcid = r[m1['Root_Cause_ID']]
        summary = r[m1['Finding_Summary']]
        file_path = r[m1['Current_Main_File']] or r[m1['Original_File']]
        line = r[m1['Current_Main_Line']] or r[m1['Original_Line']]
        pr_num = r[m1['PR_Number']]
        sev = r[m1['Current_Severity']]
        active_findings_by_rc[rcid].append({
            'fid': fid,
            'summary': summary,
            'file': file_path,
            'line': line,
            'pr': pr_num,
            'sev': sev,
            'status': st,
        })

MACRO_ISSUES = [
    {
        "index": 1,
        "title": "[Ledger 1/4] Compiler Frontend & Semantic Core Closure",
        "branch": "fix/ledger-1-compiler-semantic-core",
        "prereq": "REM-001 prerequisite satisfied by merged PR #1965 (commit `7abb364e`).",
        "blocks": "Blocks [Ledger 2/4].",
        "purpose": (
            "Consolidates and closes all remaining compiler frontend (`sm-front`), intermediate representation "
            "(`sm-ir`), and normative language/quad specification debt across 6 remediation packages (36 findings). "
            "Addresses pattern destructuring, expression positions, arithmetic coercion, effect isolation, ownership "
            "event transport, sequence equality, and specification alignment."
        ),
        "scope_summary": (
            "- `crates/sm-front/**`: tuple destructuring, promoted patterns, parser grammar, loop/array generic grammar, fx coercion, effect isolation\n"
            "- `crates/sm-ir/**`: ownership/borrow event transport, sequence equality, intrinsic shape lowering\n"
            "- `docs/spec/**`: language syntax, Logos, and quad logic formal specification alignment"
        ),
        "rem_list": ["REM-002", "REM-003", "REM-004", "REM-005", "REM-006", "REM-023"],
        "internal_deps": "`REM-004 -> REM-005`. REM-023 documentation reconciliation follows implementation truth, not leading it."
    },
    {
        "index": 2,
        "title": "[Ledger 2/4] Execution, CLI & Trust Boundary Closure",
        "branch": "fix/ledger-2-execution-cli-trust",
        "prereq": "Blocked by [Ledger 1/4] (Macro PR 1/4 must be merged first, e.g. REM-019 depends on REM-005; REM-001 satisfied by PR #1965).",
        "blocks": "Blocks [Ledger 3/4].",
        "purpose": (
            "Consolidates and closes deterministic VM execution (`sm-vm`), shared runtime core (`sm-runtime-core`), "
            "CLI sandboxing (`smc-cli`), host capability gating (`prom-cap`), integration test suites, and VM/verifier "
            "specification alignment across 13 remediation packages (82 findings)."
        ),
        "scope_summary": (
            "- `crates/sm-vm/**`: runtime ownership tracking, StoreVar synchronization, host effect buffering, streaming observation\n"
            "- `crates/sm-runtime-core/**`: API/no_std contracts, allocation boundaries, verifier specifications\n"
            "- `crates/smc-cli/**`: manifest path normalization, project root sandboxing, symlink traversal security, 7hell diagnostics\n"
            "- `crates/semantic-core/**`: specialized arithmetic safety, backend probing, bitmask primitives\n"
            "- `crates/prom-cap/**`: capability gate channel validation, audit tamper-resistance\n"
            "- `tests/**`: integration test suite quality, golden trace determinism, project model acceptance, test log sanitization\n"
            "- `docs/spec/**` & `docs/roadmap/**`: VM execution & verifier architecture alignment, compiler stage mapping"
        ),
        "rem_list": [
            "REM-007", "REM-008", "REM-009", "REM-010", "REM-011", "REM-012", "REM-013",
            "REM-018", "REM-019", "REM-020", "REM-021", "REM-024", "REM-026"
        ],
        "internal_deps": "`REM-018 <- REM-008`, `REM-019 <- REM-005 + REM-008`, `REM-020 <- REM-010`, `REM-026 <- REM-011`."
    },
    {
        "index": 3,
        "title": "[Ledger 3/4] Native UI & Workbench Closure",
        "branch": "fix/ledger-3-native-ui-workbench",
        "prereq": "Blocked by [Ledger 2/4] (Macro PR 2/4 must be merged first).",
        "blocks": "Blocks [Ledger 4/4].",
        "purpose": (
            "Consolidates and closes native platform UI windowing (`prom-ui-backend-native`), runtime sessions (`prom-ui-runtime`), "
            "presentation models (`prom-ui`), interactive demo harness, and UI envelope specifications across 5 remediation packages (64 findings)."
        ),
        "scope_summary": (
            "- `crates/prom-ui-backend-native/**`: native windowing, winit platform backends, event loop safety\n"
            "- `crates/prom-ui-runtime/**`: session lifecycle, state transitions, shell player authorization\n"
            "- `crates/prom-ui/**`: presentation models, renderer trace integration, AST slot bridge, demo harness, hit-testing\n"
            "- `docs/spec/**`: UI architecture, effect envelopes, and shell player specification alignment"
        ),
        "rem_list": ["REM-014", "REM-015", "REM-016", "REM-017", "REM-025"],
        "internal_deps": "`REM-014, REM-015 -> REM-016 -> REM-017 -> REM-025` (final spec reconciliation follows implementation truth)."
    },
    {
        "index": 4,
        "title": "[Ledger 4/4] CI, Governance & Repository-State Closure",
        "branch": "fix/ledger-4-ci-governance-status",
        "prereq": "Blocked by [Ledger 3/4] (Macro PR 3/4 must be merged first; final closure contour).",
        "blocks": "None (final remediation contour; leads to final Health Ledger zero-debt signoff).",
        "purpose": (
            "Final closure contour: reconciles repository automation scripts, local preflight tooling, and release-facing "
            "governance documentation across 2 remediation packages (34 findings) after all implementation truth in Issues 1–3 has landed."
        ),
        "scope_summary": (
            "- `scripts/**`: workspace script portability, local CI preflight, harness hygiene, exit behavior\n"
            "- `docs/roadmap/**`, `docs/status/**`, root docs: feature maturity matrix, status vocabulary, public maturity claims, governance reconciliation"
        ),
        "rem_list": ["REM-022", "REM-027"],
        "internal_deps": "`REM-022 -> REM-027` (tooling and script hygiene precedes final roadmap and maturity signoff)."
    }
]

COMMON_PROTOCOL = """## Execution Contract

> [!IMPORTANT]
> **ONE ISSUE = ONE PULL REQUEST**
> This Macro Issue maps to **EXACTLY ONE** remediation Pull Request (`{branch}`).
> The PR may contain multiple ordered checkpoints, commits, and review rounds,
> but must **never** be silently split into additional PRs.
> Any split requires explicit repository-owner authorization.

---

## Mandatory Execution Rounds

Every Macro PR follows the canonical 9-round methodology:

1. **Round 0 — Evidence Freeze**:
   - Revalidate every included finding against current `main`.
   - Record technical state: `ACTIVE_CODE_GAP`, `ALREADY_FIXED_NEEDS_EVIDENCE`, `ARCHITECTURALLY_SUPERSEDED`, `DOC_ALIGNMENT_ONLY`, `FALSE_POSITIVE_CANDIDATE`, or `UNVERIFIED`.
   - Zero production changes during Round 0.
2. **Round 1 — Tests & Reproduction**:
   - For confirmed code gaps: construct **RED regression test first**.
   - For already-fixed debt: identify or add **GREEN regression anchor**.
   - For architectural supersession: document code + spec + history evidence.
   - For doc-only findings: verify actual implementation behavior first.
3. **Round 2 — Implementation Wave**:
   - Fix only confirmed active gaps with minimal semantic footprint.
   - Reuse canonical shared validation boundaries; no drive-by cleanup or redesigns.
4. **Round 3 — First Review**:
   - Request `@codex review` and Copilot review on exact PR head.
   - Classify every thread: `REAL_DEFECT`, `DESIGN_CONFLICT`, `SPEC_DRIFT`, `FALSE_POSITIVE`, `ALREADY_ADDRESSED`.
5. **Round 4 — Repair Round**:
   - Address substantive review feedback with targeted changes and fresh test runs.
6. **Round 5 — Second Review**:
   - Mandatory reviewer re-summon after repairs until **0 unresolved substantive threads**.
7. **Round 6 — Full Qualification**:
   - Run crate tests, public API guards, boundary enforcement, clippy (`-D warnings`), formatting, hygiene, and harness check.
8. **Round 7 — Merge Readiness**:
   - Simultaneously require: all CI checks green, final review approved on exact head, 0 unresolved threads, mergeable.
9. **Round 8 — Post-Merge Reconciliation**:
   - Update derived Ledger state (`Semantic_Project_Health_Ledger.xlsx`) and regenerate `ledger_data.json` & HTML dashboard.
"""

all_mapped_rems = []
all_mapped_findings = []

for issue in MACRO_ISSUES:
    idx = issue["index"]
    title = issue["title"]
    branch = issue["branch"]
    prereq = issue["prereq"]
    blocks = issue["blocks"]
    purpose = issue["purpose"]
    scope_summary = issue["scope_summary"]
    rems = issue["rem_list"]
    int_deps = issue["internal_deps"]

    issue_rems = []
    issue_findings = []

    # Build Scope Table
    scope_rows = []
    for qid in rems:
        all_mapped_rems.append(qid)
        r_info = rem_data[qid]
        rcid = r_info["rc_id"]
        rc_info = rc_data[rcid]
        f_list = active_findings_by_rc[rcid]
        issue_findings.extend([f["fid"] for f in f_list])
        all_mapped_findings.extend([f["fid"] for f in f_list])

        scope_rows.append(
            f"| `{qid}` | `{rcid}` | `{r_info['area']}` | {rc_info['title']} | {len(f_list)} | {r_info['priority']} | {r_info['risk_tier']} | `{r_info['deps']}` |"
        )

    scope_table_str = "\n".join(scope_rows)

    # Build Finding Inventory details
    inventory_blocks = []
    for qid in rems:
        r_info = rem_data[qid]
        rcid = r_info["rc_id"]
        rc_info = rc_data[rcid]
        f_list = active_findings_by_rc[rcid]

        finding_items = []
        for f in f_list:
            finding_items.append(
                f"- [ ] **`{f['fid']}`** ({f['sev']}): {f['summary']} (origin: PR #{f['pr']}, file: `{f['file']}`)"
            )
        items_str = "\n".join(finding_items)

        inventory_blocks.append(
            f"<details>\n<summary><b>{qid} / {rcid} — {len(f_list)} findings</b> ({rc_info['title']})</summary>\n\n"
            f"{items_str}\n\n</details>"
        )

    inventory_str = "\n\n".join(inventory_blocks)

    body = f"""# {title}

**Target Branch**: `{branch}`  
**Prerequisite**: {prereq}  
**Sequencing**: {blocks}  
**Total Scope**: **{len(rems)} REM Packages** | **{len(issue_findings)} Active Findings**

---

## 1. Architectural Purpose & Scope

{purpose}

### Subsystem Boundaries
{scope_summary}

### Internal Package Dependencies
{int_deps}

---

## 2. Package Scope & Metadata Table

| REM_ID | Root_Cause_ID | Area | Theme | Active Findings | Priority | Risk Tier | Dependencies |
| :--- | :--- | :--- | :--- | :---: | :---: | :---: | :--- |
{scope_table_str}

---

## 3. Exact Finding Inventory ({len(issue_findings)} Findings)

{inventory_str}

---

{COMMON_PROTOCOL.format(branch=branch)}
"""

    out_path = os.path.join(SCRATCH_DIR, f"issue_{idx}.md")
    with open(out_path, "w", encoding="utf-8") as f:
        f.write(body)
    print(f"Generated {out_path} ({len(rems)} REM, {len(issue_findings)} findings)")

print("\n--- PREFLIGHT VALIDATION ---")
print(f"REM-001 reconciled: YES (PR #1965, commit 7abb364e)")
print(f"Remaining REM:       {len(all_mapped_rems)} / 26")
print(f"Remaining findings:  {len(all_mapped_findings)} / 216")
print(f"Duplicate REM:       {len(all_mapped_rems) - len(set(all_mapped_rems))}")
print(f"Duplicate Finding:   {len(all_mapped_findings) - len(set(all_mapped_findings))}")
print(f"Missing REM:         {26 - len(set(all_mapped_rems))}")
print(f"Missing Finding:     {216 - len(set(all_mapped_findings))}")

assert len(all_mapped_rems) == 26
assert len(set(all_mapped_rems)) == 26
assert len(all_mapped_findings) == 216
assert len(set(all_mapped_findings)) == 216
assert "REM-001" not in all_mapped_rems
print("\n>>> ALL INVARIANTS SATISFIED! Ready to create GitHub Issues. <<<")
