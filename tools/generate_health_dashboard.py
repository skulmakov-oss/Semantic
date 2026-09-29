#!/usr/bin/env python3
"""
Semantic Project Health Ledger Dashboard Generator
==================================================
Deterministic, fail-closed generator that transforms the frozen
`reports/security/Semantic_Project_Health_Ledger.xlsx` into:
1. `reports/security/dashboard/ledger_data.json`
2. `reports/security/dashboard/ledger_data.js` (for zero-config offline file:// usage)
3. `reports/security/Semantic_Project_Health_Dashboard.html` (self-contained single-file bundle)

This script validates all canonical totals and referential integrity before emitting
any output. If any assertion fails, generation aborts immediately.
"""

import os
import sys
import json
import re
from datetime import datetime, timezone
from collections import defaultdict, Counter
import openpyxl

REPO_ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
LEDGER_PATH = os.path.join(REPO_ROOT, "reports", "security", "Semantic_Project_Health_Ledger.xlsx")
DASHBOARD_DIR = os.path.join(REPO_ROOT, "reports", "security", "dashboard")
JSON_OUT = os.path.join(DASHBOARD_DIR, "ledger_data.json")
JS_DATA_OUT = os.path.join(DASHBOARD_DIR, "ledger_data.js")
SINGLE_FILE_OUT = os.path.join(REPO_ROOT, "reports", "security", "Semantic_Project_Health_Dashboard.html")
HTML_TEMPLATE = os.path.join(DASHBOARD_DIR, "index.html")
CSS_FILE = os.path.join(DASHBOARD_DIR, "styles.css")
JS_APP_FILE = os.path.join(DASHBOARD_DIR, "app.js")

EXPECTED_TOTALS = {
    "total_prs": 1477,
    "merged_prs": 1372,
    "unmerged_prs": 105,
    "review_threads": 819,
    "unresolved_threads": 375,
    "candidate_findings": 813,
    "still_present": 176,
    "partially_fixed": 4,
    "fixed_later": 481,
    "obsolete": 47,
    "false_positive": 97,
    "unverified": 8,
    "active_debt": 180,
    "root_causes": 29,
    "active_root_causes": 20,
    "resolved_root_causes": 9,
    "remediation_packages": 27,
    "credential_exposures": 0,
    "privacy_path_leaks": 119,
    "codeql_alerts": 1,
    "total_security_rows": 120,
}

def load_sheet_records(ws):
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(c).strip() if c is not None else f"col_{i}" for i, c in enumerate(rows[0])]
    records = []
    for r in rows[1:]:
        rec = {}
        for h, v in zip(headers, r):
            if isinstance(v, (datetime,)):
                v = v.strftime("%Y-%m-%d")
            rec[h] = v
        records.append(rec)
    return records

def validate_and_extract():
    if not os.path.exists(LEDGER_PATH):
        sys.exit(f"FATAL: Ledger file not found at {LEDGER_PATH}")

    print(f"Loading ledger from: {LEDGER_PATH}...")
    wb = openpyxl.load_workbook(LEDGER_PATH, data_only=True)
    required_sheets = [
        "00_Dashboard", "01_Findings", "02_Root_Causes", "03_PR_Coverage",
        "04_Review_Threads", "05_SSF_Mapping", "06_Security_Privacy",
        "07_Remediation_Queue", "08_Dictionaries", "09_Metadata"
    ]
    for s in required_sheets:
        if s not in wb.sheetnames:
            sys.exit(f"FATAL: Missing required sheet '{s}' in ledger!")

    # 1. Findings
    findings = load_sheet_records(wb["01_Findings"])
    if len(findings) != EXPECTED_TOTALS["candidate_findings"]:
        sys.exit(f"ASSERTION FAILED: Expected {EXPECTED_TOTALS['candidate_findings']} findings, got {len(findings)}")

    status_counts = defaultdict(int)
    sev_counts = defaultdict(int)
    active_sev_counts = defaultdict(int)
    area_counts = defaultdict(int)
    active_area_counts = defaultdict(int)
    type_counts = defaultdict(int)
    sec_class_counts = defaultdict(int)
    ssf_rel_counts = defaultdict(int)

    fnd_id_set = set()
    for f in findings:
        fid = f["Finding_ID"]
        fnd_id_set.add(fid)
        st = f["Current_Status"]
        sev = f["Current_Severity"]
        area = f["Area"]
        ft = f.get("Finding_Type") or "UNKNOWN"
        sc = f.get("Security_Class") or "NONE"
        ssf = f.get("SSF_Relation") or "NO_RELATION"

        status_counts[st] += 1
        sev_counts[sev] += 1
        area_counts[area] += 1
        type_counts[ft] += 1
        sec_class_counts[sc] += 1
        ssf_rel_counts[ssf] += 1

        if st in ("STILL_PRESENT", "PARTIALLY_FIXED"):
            active_sev_counts[sev] += 1
            active_area_counts[area] += 1

    active_debt = status_counts["STILL_PRESENT"] + status_counts["PARTIALLY_FIXED"]
    assert status_counts["STILL_PRESENT"] == EXPECTED_TOTALS["still_present"], f"still_present mismatch: {status_counts['STILL_PRESENT']}"
    assert status_counts["PARTIALLY_FIXED"] == EXPECTED_TOTALS["partially_fixed"], f"partially_fixed mismatch: {status_counts['PARTIALLY_FIXED']}"
    assert status_counts["FIXED_LATER"] == EXPECTED_TOTALS["fixed_later"], f"fixed_later mismatch: {status_counts['FIXED_LATER']}"
    assert status_counts["OBSOLETE"] == EXPECTED_TOTALS["obsolete"], f"obsolete mismatch: {status_counts['OBSOLETE']}"
    assert status_counts["FALSE_POSITIVE"] == EXPECTED_TOTALS["false_positive"], f"false_positive mismatch: {status_counts['FALSE_POSITIVE']}"
    assert status_counts["UNVERIFIED"] == EXPECTED_TOTALS["unverified"], f"unverified mismatch: {status_counts['UNVERIFIED']}"
    assert active_debt == EXPECTED_TOTALS["active_debt"], f"active_debt mismatch: {active_debt}"

    # 2. Root Causes
    root_causes = load_sheet_records(wb["02_Root_Causes"])
    assert len(root_causes) == EXPECTED_TOTALS["root_causes"], f"root_causes mismatch: {len(root_causes)}"
    rc_id_set = {r["Root_Cause_ID"] for r in root_causes}
    active_rc_count = sum(1 for r in root_causes if r.get("Status") == "ACTIVE")
    resolved_rc_count = sum(1 for r in root_causes if r.get("Status") == "RESOLVED")
    assert active_rc_count == EXPECTED_TOTALS["active_root_causes"], f"active_rc mismatch: {active_rc_count}"
    assert resolved_rc_count == EXPECTED_TOTALS["resolved_root_causes"], f"resolved_rc mismatch: {resolved_rc_count}"

    # 3. PR Coverage
    pr_coverage = load_sheet_records(wb["03_PR_Coverage"])
    assert len(pr_coverage) == EXPECTED_TOTALS["total_prs"], f"total_prs mismatch: {len(pr_coverage)}"
    merged_prs = sum(1 for p in pr_coverage if p.get("Merged") == "YES")
    unmerged_prs = sum(1 for p in pr_coverage if p.get("Merged") == "NO")
    assert merged_prs == EXPECTED_TOTALS["merged_prs"], f"merged_prs mismatch: {merged_prs}"
    assert unmerged_prs == EXPECTED_TOTALS["unmerged_prs"], f"unmerged_prs mismatch: {unmerged_prs}"
    pr_number_set = {p["PR_Number"] for p in pr_coverage}

    # 4. Review Threads
    review_threads = load_sheet_records(wb["04_Review_Threads"])
    assert len(review_threads) == EXPECTED_TOTALS["review_threads"], f"review_threads mismatch: {len(review_threads)}"
    unres_threads = sum(1 for t in review_threads if t.get("Thread_Resolved") == "NO")
    assert unres_threads == EXPECTED_TOTALS["unresolved_threads"], f"unresolved_threads mismatch: {unres_threads}"

    # 5. SSF Mapping
    ssf_mapping = load_sheet_records(wb["05_SSF_Mapping"])
    assert len(ssf_mapping) == EXPECTED_TOTALS["candidate_findings"], f"ssf_mapping mismatch: {len(ssf_mapping)}"

    # 6. Security & Privacy
    sec_privacy = load_sheet_records(wb["06_Security_Privacy"])
    assert len(sec_privacy) == EXPECTED_TOTALS["total_security_rows"], f"total_security_rows mismatch: {len(sec_privacy)}"
    cred_exposures = sum(1 for s in sec_privacy if s.get("Credential_Exposure") == "YES")
    privacy_exposures = sum(1 for s in sec_privacy if s.get("Privacy_Exposure") == "YES")
    codeql_alerts = sum(1 for s in sec_privacy if s.get("Source_Type") == "CODEQL_ALERT")
    assert cred_exposures == EXPECTED_TOTALS["credential_exposures"], f"credential_exposures mismatch: {cred_exposures}"
    assert privacy_exposures == EXPECTED_TOTALS["privacy_path_leaks"], f"privacy_path_leaks mismatch: {privacy_exposures}"
    assert codeql_alerts == EXPECTED_TOTALS["codeql_alerts"], f"codeql_alerts mismatch: {codeql_alerts}"

    # Zero PR check in security
    zero_sec_prs = sum(1 for s in sec_privacy if s.get("PR_Number") in (0, "0", None, "", "UNKNOWN"))
    assert zero_sec_prs == 0, f"Found {zero_sec_prs} security records with PR_Number = 0"

    # 7. Remediation Queue
    rem_queue = load_sheet_records(wb["07_Remediation_Queue"])
    assert len(rem_queue) == EXPECTED_TOTALS["remediation_packages"], f"remediation_packages mismatch: {len(rem_queue)}"
    rem_id_set = {r["Queue_ID"] for r in rem_queue}
    rem_findings_total = sum(int(r.get("Finding_Count") or 0) for r in rem_queue)
    assert rem_findings_total == EXPECTED_TOTALS["active_debt"], f"rem_findings_total mismatch: {rem_findings_total}"

    # 8. Dictionaries & Metadata
    dict_records = load_sheet_records(wb["08_Dictionaries"])
    dictionaries = defaultdict(list)
    for row in dict_records:
        for k, v in row.items():
            if v and str(v).strip():
                dictionaries[k].append(str(v).strip())

    meta_records = load_sheet_records(wb["09_Metadata"])
    metadata = {}
    for r in meta_records:
        k = r.get("Metadata Key")
        v = r.get("Metadata Value")
        if k:
            metadata[k] = v

    # 9. Referential Integrity Validation
    for f in findings:
        rc = f.get("Root_Cause_ID")
        if rc and rc not in rc_id_set:
            sys.exit(f"REFERENTIAL INTEGRITY ERROR: Finding {f['Finding_ID']} references invalid Root Cause '{rc}'")
        rg = f.get("Remediation_Group")
        st = f.get("Current_Status")
        if st in ("STILL_PRESENT", "PARTIALLY_FIXED"):
            if not rg or rg not in rem_id_set:
                sys.exit(f"REFERENTIAL INTEGRITY ERROR: Active Finding {f['Finding_ID']} references invalid Remediation Group '{rg}'")
        pnum = f.get("PR_Number")
        if pnum and pnum not in pr_number_set:
            sys.exit(f"REFERENTIAL INTEGRITY ERROR: Finding {f['Finding_ID']} references unknown PR #{pnum}")

    for t in review_threads:
        pnum = t.get("PR_Number")
        if pnum and pnum not in pr_number_set:
            sys.exit(f"REFERENTIAL INTEGRITY ERROR: Thread {t.get('Thread_ID')} references unknown PR #{pnum}")

    for s in sec_privacy:
        pnum = s.get("PR_Number")
        if pnum and pnum not in pr_number_set:
            sys.exit(f"REFERENTIAL INTEGRITY ERROR: Security record {s.get('Security_ID')} references unknown PR #{pnum}")

    for r in rem_queue:
        rc = r.get("Root_Cause_ID")
        if rc and rc not in rc_id_set:
            sys.exit(f"REFERENTIAL INTEGRITY ERROR: Remediation queue {r.get('Queue_ID')} references unknown Root Cause '{rc}'")
        deps = r.get("Dependencies")
        if deps and deps != "None":
            dep_list = [d.strip() for d in str(deps).split(",") if d.strip()]
            for d in dep_list:
                if d not in rem_id_set:
                    sys.exit(f"REFERENTIAL INTEGRITY ERROR: Remediation queue {r.get('Queue_ID')} references unknown dependency '{d}'")

    print("ALL CANONICAL ASSERTIONS AND REFERENTIAL INTEGRITY CHECKS PASSED (100% VALID)!")

    # 10. Precomputed chart data
    # PR Era buckets
    era_buckets = [
        {"era": "PR 1 - 250", "min": 1, "max": 250, "total": 0, "active": 0},
        {"era": "PR 251 - 500", "min": 251, "max": 500, "total": 0, "active": 0},
        {"era": "PR 501 - 750", "min": 501, "max": 750, "total": 0, "active": 0},
        {"era": "PR 751 - 1000", "min": 751, "max": 1000, "total": 0, "active": 0},
        {"era": "PR 1001 - 1250", "min": 1001, "max": 1250, "total": 0, "active": 0},
        {"era": "PR 1251 - 1500", "min": 1251, "max": 1500, "total": 0, "active": 0},
    ]
    for f in findings:
        pnum = f.get("PR_Number") or 0
        is_act = f.get("Current_Status") in ("STILL_PRESENT", "PARTIALLY_FIXED")
        for b in era_buckets:
            if b["min"] <= pnum <= b["max"]:
                b["total"] += 1
                if is_act:
                    b["active"] += 1
                break

    # Subsystem active debt sorted
    subsystems_sorted = sorted(active_area_counts.items(), key=lambda x: x[1], reverse=True)
    rem_areas_sorted = sorted(Counter([r.get("Area") for r in rem_queue]).items(), key=lambda x: x[1], reverse=True)

    summary = {
        "baseline_sha": "e41ff311",
        "baseline_branch": "origin/main",
        "ledger_version": "v1.0-frozen",
        "generated_at": datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"),
        "total_prs": EXPECTED_TOTALS["total_prs"],
        "merged_prs": EXPECTED_TOTALS["merged_prs"],
        "unmerged_prs": EXPECTED_TOTALS["unmerged_prs"],
        "review_threads": EXPECTED_TOTALS["review_threads"],
        "unresolved_threads": EXPECTED_TOTALS["unresolved_threads"],
        "resolved_threads": EXPECTED_TOTALS["review_threads"] - EXPECTED_TOTALS["unresolved_threads"],
        "candidate_findings": EXPECTED_TOTALS["candidate_findings"],
        "active_debt": EXPECTED_TOTALS["active_debt"],
        "inactive_findings": EXPECTED_TOTALS["candidate_findings"] - EXPECTED_TOTALS["active_debt"],
        "still_present": status_counts["STILL_PRESENT"],
        "partially_fixed": status_counts["PARTIALLY_FIXED"],
        "fixed_later": status_counts["FIXED_LATER"],
        "obsolete": status_counts["OBSOLETE"],
        "false_positive": status_counts["FALSE_POSITIVE"],
        "unverified": status_counts["UNVERIFIED"],
        "root_causes": EXPECTED_TOTALS["root_causes"],
        "active_root_causes": EXPECTED_TOTALS["active_root_causes"],
        "resolved_root_causes": EXPECTED_TOTALS["resolved_root_causes"],
        "remediation_packages": EXPECTED_TOTALS["remediation_packages"],
        "credential_exposures": 0,
        "privacy_path_leaks": EXPECTED_TOTALS["privacy_path_leaks"],
        "codeql_alerts": EXPECTED_TOTALS["codeql_alerts"],
        "affected_privacy_prs": len({s["PR_Number"] for s in sec_privacy if s.get("Privacy_Exposure") == "YES"}),
        "security_class_findings": sum(v for k, v in sec_class_counts.items() if k != "NONE"),
        "status_distribution": dict(status_counts),
        "severity_distribution": dict(sev_counts),
        "active_severity_distribution": dict(active_sev_counts),
        "active_area_distribution": dict(subsystems_sorted),
        "rem_area_distribution": dict(rem_areas_sorted),
        "ssf_relation_distribution": dict(ssf_rel_counts),
        "pr_era_buckets": era_buckets,
    }

    full_dataset = {
        "summary": summary,
        "findings": findings,
        "root_causes": root_causes,
        "remediation_queue": rem_queue,
        "pr_coverage": pr_coverage,
        "review_threads": review_threads,
        "ssf_mapping": ssf_mapping,
        "security_privacy": sec_privacy,
        "dictionaries": dict(dictionaries),
        "metadata": metadata,
    }

    return full_dataset

from collections import Counter

def main():
    os.makedirs(DASHBOARD_DIR, exist_ok=True)
    dataset = validate_and_extract()

    # 1. Write JSON
    print(f"Writing dataset to {JSON_OUT}...")
    with open(JSON_OUT, "w", encoding="utf-8") as f:
        json.dump(dataset, f, indent=2, ensure_ascii=False)
    print(f"Wrote {os.path.getsize(JSON_OUT):,} bytes to {JSON_OUT}.")

    # 2. Write JS wrapper for zero-config file:// offline viewing
    print(f"Writing JS wrapper to {JS_DATA_OUT}...")
    with open(JS_DATA_OUT, "w", encoding="utf-8") as f:
        f.write("// Canonical Ledger Data generated by tools/generate_health_dashboard.py\n")
        f.write("window.SEMANTIC_LEDGER_DATA = ")
        json.dump(dataset, f, indent=2, ensure_ascii=False)
        f.write(";\n")
    print(f"Wrote {os.path.getsize(JS_DATA_OUT):,} bytes to {JS_DATA_OUT}.")

    # 3. Generate Single-File Self-Contained Bundle if template, css, js exist
    if os.path.exists(HTML_TEMPLATE) and os.path.exists(CSS_FILE) and os.path.exists(JS_APP_FILE):
        print(f"Generating single-file self-contained dashboard at {SINGLE_FILE_OUT}...")
        with open(HTML_TEMPLATE, "r", encoding="utf-8") as f:
            html = f.read()
        with open(CSS_FILE, "r", encoding="utf-8") as f:
            css = f.read()
        with open(JS_APP_FILE, "r", encoding="utf-8") as f:
            js_app = f.read()

        # Inlining CSS
        html = re.sub(
            r'<link\s+rel="stylesheet"\s+href="styles\.css"[^>]*>',
            f"<style>\n{css}\n</style>",
            html
        )
        # Inlining JS data and App
        js_embed = (
            f"<script>\nwindow.SEMANTIC_LEDGER_DATA = {json.dumps(dataset, ensure_ascii=False)};\n</script>\n"
            f"<script>\n{js_app}\n</script>"
        )
        html = re.sub(
            r'<script\s+src="ledger_data\.js"[^>]*></script>\s*<script\s+src="app\.js"[^>]*></script>',
            js_embed,
            html
        )
        # Fallback if scripts are separated
        if '<script src="ledger_data.js"></script>' in html:
            html = html.replace(
                '<script src="ledger_data.js"></script>',
                f"<script>\nwindow.SEMANTIC_LEDGER_DATA = {json.dumps(dataset, ensure_ascii=False)};\n</script>"
            )
        if '<script src="app.js"></script>' in html:
            html = html.replace(
                '<script src="app.js"></script>',
                f"<script>\n{js_app}\n</script>"
            )

        with open(SINGLE_FILE_OUT, "w", encoding="utf-8") as f:
            f.write(html)
        print(f"Single-file distribution created at {SINGLE_FILE_OUT} ({os.path.getsize(SINGLE_FILE_OUT):,} bytes).")

    print("\nSUCCESS: All dashboard data artifacts successfully generated and verified!")

if __name__ == "__main__":
    main()
